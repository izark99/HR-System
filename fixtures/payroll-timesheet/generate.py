"""Sinh dữ liệu GIẢ cho flow nghiệm thu "Đối chiếu bảng lương – chấm công".

Chạy: python3 fixtures/payroll-timesheet/generate.py
Kết quả: các file .xlsx + expected.json trong cùng thư mục.
Toàn bộ dữ liệu là ngẫu nhiên (seed cố định), không liên quan nhân sự thật.
"""

import json
import random
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

OUT = Path(__file__).parent
PERIOD = "09/2026"
STANDARD_DAYS = Decimal("26")

# Ngưỡng mặc định dùng để tính expected.json (Admin cấu hình được trong node compare)
TOLERANCE = {"cong": Decimal("0.5"), "ot": Decimal("1"), "phep": Decimal("0"), "khong_luong": Decimal("0")}

BRANCHES = [
    # (mã, tên, số NV, header mã NV của file lương)
    ("CN01", "Hồ Chí Minh", 30, "Mã NV"),
    ("CN02", "Cần Thơ", 25, "MSNV"),  # alias khác để test so khớp tên cột
    ("CN03", "Hà Nội", 20, "Mã NV"),
]
DEPARTMENTS = ["Sản xuất", "Kho vận", "Kinh doanh", "QA", "QC", "Hành chính"]
LAST = ["Nguyễn", "Trần", "Lê", "Phạm", "Hoàng", "Võ", "Đặng", "Bùi", "Đỗ", "Huỳnh"]
MIDDLE = ["Văn", "Thị", "Minh", "Ngọc", "Hữu", "Thanh", "Quốc", "Gia"]
FIRST = ["An", "Bình", "Chi", "Dũng", "Hà", "Khoa", "Lan", "Mai", "Nam", "Phúc", "Quân", "Thảo", "Vy"]

rng = random.Random(20261026)


def d(x: float) -> Decimal:
    return Decimal(str(round(x * 2) / 2))  # làm tròn 0.5


def fake_employee(emp_no: int, branch: str) -> dict:
    phep = d(rng.choice([0, 0, 0, 1, 1, 2]))
    kl = d(rng.choice([0, 0, 0, 0, 1]))
    return {
        "ma_nv": f"NV{emp_no:05d}",
        "ho_ten": f"{rng.choice(LAST)} {rng.choice(MIDDLE)} {rng.choice(FIRST)}",
        "phong_ban": rng.choice(DEPARTMENTS),
        "chi_nhanh": branch,
        "cong": STANDARD_DAYS - phep - kl,
        "ot": d(rng.uniform(0, 20)),
        "phep": phep,
        "khong_luong": kl,
        "luong_cb": Decimal(rng.randrange(6_000_000, 30_000_000, 100_000)),
    }


def write_payroll(branch_code, branch_name, id_header, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Luong"
    ws["A1"] = f"BẢNG LƯƠNG THÁNG {PERIOD} – CHI NHÁNH {branch_name.upper()}"
    ws.merge_cells("A1:J1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "Dữ liệu giả phục vụ kiểm thử"
    headers = ["STT", id_header, "Họ và tên", "Phòng ban", "Chi nhánh", "Lương CB",
               "Công tính lương", "Giờ OT", "Ngày phép", "Nghỉ không lương"]
    ws.append([])  # dòng 3 trống → header ở dòng 4
    ws.append(headers)
    for c in ws[4]:
        c.font = Font(bold=True)
    for i, r in enumerate(rows, 1):
        ws.append([i, r["ma_nv"], r["ho_ten"], r["phong_ban"], r["chi_nhanh"], float(r["luong_cb"]),
                   float(r["cong"]), float(r["ot"]), float(r["phep"]), float(r["khong_luong"])])
    ws.append([])
    ws.append(["", "", "TỔNG CỘNG", "", "", f"=SUM(F5:F{4 + len(rows)})"])  # dòng tổng: phải bị loại khi đọc
    wb.save(OUT / f"bang_luong_{branch_code}_{PERIOD.replace('/', '')}.xlsx")


def write_timesheet(branch_code, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "ChamCong"
    ws.append(["Mã nhân viên", "Họ tên", "Chi nhánh", "Công thực tế", "OT (giờ)", "Phép (ngày)", "Không lương (ngày)"])
    for r in rows:
        ws.append([r["ma_nv"], r["ho_ten"], r["chi_nhanh"], float(r["cong"]), float(r["ot"]),
                   float(r["phep"]), float(r["khong_luong"])])
    wb.save(OUT / f"cham_cong_{branch_code}_{PERIOD.replace('/', '')}.xlsx")


def main():
    emp_no = 1
    payroll_all, timesheet_all = [], []
    for code, name, n, id_header in BRANCHES:
        payroll = [fake_employee(emp_no + i, code) for i in range(n)]
        emp_no += n
        timesheet = [dict(r) for r in payroll]

        # Tạo sai lệch có chủ đích trên file chấm công
        idx = rng.sample(range(n), 6)
        timesheet[idx[0]]["cong"] -= Decimal("0.5")   # lệch trong ngưỡng → khớp
        timesheet[idx[1]]["cong"] -= Decimal("2")     # lệch công vượt ngưỡng
        timesheet[idx[2]]["ot"] += Decimal("4")       # lệch OT vượt ngưỡng
        timesheet[idx[3]]["phep"] += Decimal("1")     # lệch phép
        timesheet[idx[3]]["cong"] -= Decimal("1")     # (đồng thời lệch công)
        only_payroll = timesheet.pop(idx[4])          # NV chỉ có trong lương
        extra = fake_employee(90000 + emp_no, code)   # NV chỉ có trong chấm công
        timesheet.append(extra)

        write_payroll(code, name, id_header, payroll)
        write_timesheet(code, timesheet)
        payroll_all += payroll
        timesheet_all += timesheet
        del only_payroll

    # Tính kết quả kỳ vọng
    p = {r["ma_nv"]: r for r in payroll_all}
    t = {r["ma_nv"]: r for r in timesheet_all}
    mismatched, matched = [], []
    for k in sorted(p.keys() & t.keys()):
        diffs = {m: str(t[k][m] - p[k][m]) for m in TOLERANCE if abs(t[k][m] - p[k][m]) > TOLERANCE[m]}
        (mismatched if diffs else matched).append({"ma_nv": k, **({"diffs": diffs} if diffs else {})})
    expected = {
        "period": PERIOD,
        "tolerance": {k: str(v) for k, v in TOLERANCE.items()},
        "counts": {
            "payroll_rows": len(payroll_all),
            "timesheet_rows": len(timesheet_all),
            "matched": len(matched),
            "mismatched": len(mismatched),
            "only_in_payroll": len(p.keys() - t.keys()),
            "only_in_timesheet": len(t.keys() - p.keys()),
        },
        "mismatched": mismatched,
        "only_in_payroll": sorted(p.keys() - t.keys()),
        "only_in_timesheet": sorted(t.keys() - p.keys()),
    }
    (OUT / "expected.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(expected["counts"], indent=2))


if __name__ == "__main__":
    main()
