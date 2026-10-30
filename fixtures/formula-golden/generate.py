"""Sinh file golden cho node `transform.formula` (spec §12.2, AC-B4-30).

Chạy: python3 fixtures/formula-golden/generate.py
Kết quả (cùng thư mục):
  - golden.xlsx : sheet "cases", mỗi dòng 1 case; cột G là CÔNG THỨC Excel thật (chưa có giá trị).
  - cases.json  : cùng danh sách case, kèm kiểu dữ liệu đầu vào và kết quả mong đợi của engine.

Harold mở golden.xlsx bằng Excel desktop, bật Enable Editing, nhấn Ctrl+Alt+F9 (tính lại toàn bộ),
rồi lưu thành golden-excel.xlsx trong thư mục này. Test `formula-golden.test.ts` đọc giá trị Excel
đã lưu (cached values) ở cột G và so với kết quả engine.

Toàn bộ dữ liệu là dữ liệu giả, không liên quan nhân sự thật.
"""

import datetime as dt
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).parent

# Hàm Excel 2019+ phải có tiền tố _xlfn. trong file, nếu không Excel báo #NAME?.
XLFN = ["CONCAT", "TEXTJOIN", "IFS", "SWITCH"]

# Đầu vào: (kiểu engine, giá trị); giá trị None = ô trống.
#   n = decimal (chuỗi thập phân), s = string, d = date "YYYY-MM-DD", b = boolean
# Kết quả mong đợi của engine: ("number", "2.68") | ("string", "...") | ("boolean", True)
#   | ("date", "2026-02-28") | ("error", "#DIV/0!") | ("blank", None)
# `excel`: khác biệt có chủ đích — giá trị Excel mong đợi khi khác engine (kèm `note`).
# `confirm`: hành vi Excel chưa chắc chắn, cần golden xác nhận.


def n(v):
    return ("n", v)


def s(v):
    return ("s", v)


def d(v):
    return ("d", v)


def b(v):
    return ("b", v)


BLANK_N = ("n", None)
BLANK_S = ("s", None)

CASES = []


def case(fn, formula, inputs, expect, **extra):
    CASES.append({"fn": fn, "formula": formula, "inputs": inputs, "expect": expect, **extra})


# ---- Toán tử & ô trống ---------------------------------------------------------------------
case("OP", "=[a]+1", {"a": BLANK_N}, ("number", "1"))
case("OP", '=[a]&"x"', {"a": BLANK_S}, ("string", "x"))
case("OP", '=[a]=""', {"a": BLANK_N}, ("boolean", True))
case("OP", "=[a]=0", {"a": BLANK_N}, ("boolean", True))
case("OP", "=[a]", {"a": BLANK_N}, ("blank", None), excel=("number", "0"),
     note="Engine: tham chiếu cột trần với ô trống → trống (MEMORY 2026-10-30); Excel trả 0.")
case("OP", "=0.1+0.2", {}, ("number", "0.3"))
case("OP", "=1/3", {}, ("number", "0.3333333333"), excel=("number", "0.333333333333333"),
     note="Engine: DECIMAL(38,10) làm tròn 10 chữ số; Excel 15 chữ số có nghĩa.")
case("OP", "=[a]*1.1", {"a": n("1234567890123.25")}, ("number", "1358024679135.575"),
     excel=("number", "1358024679135.58"),
     note="Engine chính xác tuyệt đối; Excel chỉ giữ 15 chữ số có nghĩa.")
case("OP", "=[a]/[b]", {"a": n("10"), "b": n("0")}, ("error", "#DIV/0!"))
case("OP", "=[a]/[b]", {"a": n("-7"), "b": n("2")}, ("number", "-3.5"))
case("OP", '="12"+1', {}, ("number", "13"))
case("OP", '="abc"+1', {}, ("error", "#VALUE!"))
case("OP", "=TRUE+1", {}, ("number", "2"))
case("OP", "=-2^2", {}, ("number", "4"))
case("OP", "=2^10", {}, ("number", "1024"))
case("OP", "=2^-2", {}, ("number", "0.25"))
case("OP", "=10^30*10^30", {}, ("error", "#NUM!"), excel=("number", "1e60"),
     note="Engine: vượt DECIMAL(38,10) → #NUM!; Excel (double) vẫn tính được.")
case("OP", "=50%", {}, ("number", "0.5"))
case("OP", "=[a]%", {"a": n("-12.5")}, ("number", "-0.125"))
case("OP", '="abc"="ABC"', {}, ("boolean", True))
case("OP", '="a"<"B"', {}, ("boolean", True))
case("OP", '=1<"a"', {}, ("boolean", True))
case("OP", "=[a]<>[b]", {"a": s("Nguyễn"), "b": s("NGUYỄN")}, ("boolean", False))
case("OP", "=[a]+30", {"a": d("2024-02-15")}, ("date", "2024-03-16"))
case("OP", "=[a]-[b]", {"a": d("2026-03-01"), "b": d("2024-02-28")}, ("number", "732"))
case("OP", '="Mã: "&[a]', {"a": n("1.50")}, ("string", "Mã: 1.5"))
case("OP", "=[a]&[b]", {"a": b(True), "b": n("-3")}, ("string", "TRUE-3"))

# ---- Logic ---------------------------------------------------------------------------------
case("IF", '=IF([a]>0,"Dương","Âm")', {"a": n("5")}, ("string", "Dương"))
case("IF", '=IF([a]>0,"Dương","Âm")', {"a": n("-3")}, ("string", "Âm"))
case("IF", "=IF([a]>0,[a]*2,[a]-1)", {"a": BLANK_N}, ("number", "-1"))
case("IF", "=IF([a]/[b]>1,1,2)", {"a": n("1"), "b": n("0")}, ("error", "#DIV/0!"))
case("IF", "=IF(TRUE,[a],0)", {"a": BLANK_N}, ("blank", None), excel=("number", "0"),
     note="Engine: nhánh IF trả về tham chiếu cột trần với ô trống → trống; Excel trả 0.")
case("IF", "=IF([a]>0,1/0,5)", {"a": n("-1")}, ("number", "5"))
case("IFS", '=IFS([a]>=8,"Giỏi",[a]>=5,"Khá",TRUE,"Yếu")', {"a": n("9")}, ("string", "Giỏi"))
case("IFS", '=IFS([a]>=8,"Giỏi",[a]>=5,"Khá",TRUE,"Yếu")', {"a": n("6.5")}, ("string", "Khá"))
case("IFS", '=IFS([a]>=8,"Giỏi",[a]>=5,"Khá",TRUE,"Yếu")', {"a": BLANK_N}, ("string", "Yếu"))
case("IFS", "=IFS([a]>100,1)", {"a": n("1")}, ("error", "#N/A"))
case("AND", "=AND([a]>0,[b]>0)", {"a": n("1"), "b": n("2")}, ("boolean", True))
case("AND", "=AND([a]>0,[b]>0)", {"a": n("1"), "b": n("-2")}, ("boolean", False))
case("AND", "=AND([a]>0,[b]>0)", {"a": BLANK_N, "b": n("2")}, ("boolean", False))
case("AND", "=AND([a],1)", {"a": BLANK_N}, ("boolean", True), confirm="AND bỏ qua ô trống trong tham chiếu")
case("OR", "=OR([a]>0,[b]>0)", {"a": n("-1"), "b": n("2")}, ("boolean", True))
case("OR", "=OR([a]>0,[b]>0)", {"a": n("-1"), "b": n("-2")}, ("boolean", False))
case("OR", '=OR([a]="x",[b]=1)', {"a": BLANK_S, "b": BLANK_N}, ("boolean", False))
case("OR", "=OR(1/0,TRUE)", {}, ("error", "#DIV/0!"))
case("NOT", "=NOT([a]>5)", {"a": n("7")}, ("boolean", False))
case("NOT", "=NOT([a])", {"a": BLANK_N}, ("boolean", True))
case("NOT", "=NOT([a])", {"a": n("-2")}, ("boolean", False))
case("NOT", '=NOT("abc")', {}, ("error", "#VALUE!"))
case("SWITCH", '=SWITCH([a],1,"Một",2,"Hai","Khác")', {"a": n("1")}, ("string", "Một"))
case("SWITCH", '=SWITCH([a],1,"Một",2,"Hai","Khác")', {"a": n("3")}, ("string", "Khác"))
case("SWITCH", '=SWITCH([a],"x",1)', {"a": s("y")}, ("error", "#N/A"))
case("SWITCH", '=SWITCH([a],"hà nội","HN","SG")', {"a": s("Hà Nội")}, ("string", "HN"))
case("SWITCH", '=SWITCH([a],0,"Không","Khác")', {"a": BLANK_N}, ("string", "Không"),
     confirm="SWITCH so ô trống với 0")
case("IFERROR", "=IFERROR([a]/[b],0)", {"a": n("1"), "b": n("0")}, ("number", "0"))
case("IFERROR", "=IFERROR([a]/[b],0)", {"a": n("6"), "b": n("3")}, ("number", "2"))
case("IFERROR", "=IFERROR(VALUE([a]),-1)", {"a": s("abc")}, ("number", "-1"))
case("IFERROR", '=IFERROR(SWITCH([a],"x",1),2)', {"a": s("y")}, ("number", "2"))
case("ISBLANK", "=ISBLANK([a])", {"a": BLANK_N}, ("boolean", True))
case("ISBLANK", "=ISBLANK([a])", {"a": n("0")}, ("boolean", False))
case("ISBLANK", "=ISBLANK([a])", {"a": s("Nguyễn")}, ("boolean", False))
case("ISNUMBER", "=ISNUMBER([a])", {"a": n("5")}, ("boolean", True))
case("ISNUMBER", "=ISNUMBER([a])", {"a": s("5")}, ("boolean", False))
case("ISNUMBER", "=ISNUMBER([a])", {"a": BLANK_N}, ("boolean", False))
case("ISNUMBER", "=ISNUMBER([a])", {"a": d("2024-02-29")}, ("boolean", True))

# ---- Toán ----------------------------------------------------------------------------------
case("ROUND", "=ROUND(2.675,2)", {}, ("number", "2.68"))
case("ROUND", "=ROUND(-2.5,0)", {}, ("number", "-3"))
case("ROUND", "=ROUND([a],-2)", {"a": n("1234.5678")}, ("number", "1200"))
case("ROUND", "=ROUND([a],1)", {"a": BLANK_N}, ("number", "0"))
case("ROUND", "=ROUND([a],1)", {"a": s("abc")}, ("error", "#VALUE!"))
case("ROUND", "=ROUND(1.2345,2.9)", {}, ("number", "1.23"))
case("ROUNDUP", "=ROUNDUP(-1.21,1)", {}, ("number", "-1.3"))
case("ROUNDUP", "=ROUNDUP(1.201,2)", {}, ("number", "1.21"))
case("ROUNDUP", "=ROUNDUP([a],-1)", {"a": n("5")}, ("number", "10"))
case("ROUNDDOWN", "=ROUNDDOWN(1.29,1)", {}, ("number", "1.2"))
case("ROUNDDOWN", "=ROUNDDOWN(-1.29,1)", {}, ("number", "-1.2"))
case("ROUNDDOWN", "=ROUNDDOWN([a],-2)", {"a": n("999")}, ("number", "900"))
case("INT", "=INT(-1.5)", {}, ("number", "-2"))
case("INT", "=INT([a])", {"a": n("2.9")}, ("number", "2"))
case("INT", "=INT([a])", {"a": BLANK_N}, ("number", "0"))
case("ABS", "=ABS([a])", {"a": n("-3.5")}, ("number", "3.5"))
case("ABS", "=ABS([a])", {"a": BLANK_N}, ("number", "0"))
case("ABS", '=ABS("abc")', {}, ("error", "#VALUE!"))
case("MIN", "=MIN([a],[b],[c])", {"a": n("3"), "b": n("-1"), "c": n("2")}, ("number", "-1"))
case("MIN", "=MIN([a],[b],[c])", {"a": BLANK_N, "b": n("2"), "c": n("5")}, ("number", "2"))
case("MIN", "=MIN([a],[b])", {"a": s("abc"), "b": n("7")}, ("number", "7"))
case("MIN", "=MIN([a])", {"a": BLANK_N}, ("number", "0"))
case("MAX", "=MAX([a],[b],[c])", {"a": n("3"), "b": n("-1"), "c": n("2")}, ("number", "3"))
case("MAX", "=MAX([a],[b])", {"a": BLANK_N, "b": n("-5")}, ("number", "-5"))
case("MAX", "=MAX([a],[b])", {"a": s("abc"), "b": n("-7")}, ("number", "-7"))
case("MOD", "=MOD(-3,2)", {}, ("number", "1"))
case("MOD", "=MOD(3,-2)", {}, ("number", "-1"))
case("MOD", "=MOD([a],[b])", {"a": n("5"), "b": BLANK_N}, ("error", "#DIV/0!"))
case("MOD", "=MOD(5.5,2)", {}, ("number", "1.5"))
case("SUM", "=SUM([a],[b])", {"a": BLANK_N, "b": n("2")}, ("number", "2"))
case("SUM", "=SUM([a],[b])", {"a": n("1.1"), "b": n("2.2")}, ("number", "3.3"))
case("SUM", "=SUM([a],[b])", {"a": s("abc"), "b": n("-5")}, ("number", "-5"))
case("SUM", '=SUM("abc",1)', {}, ("error", "#VALUE!"))

# ---- Chuỗi ---------------------------------------------------------------------------------
case("LEFT", "=LEFT([a],3)", {"a": s("Nguyễn")}, ("string", "Ngu"))
case("LEFT", "=LEFT([a])", {"a": s("abc")}, ("string", "a"))
case("LEFT", "=LEFT([a],2)", {"a": BLANK_S}, ("blank", None), excel=("string", ""),
     note="Engine: chuỗi rỗng là ô trống (spec §12.1).")
case("LEFT", "=LEFT([a],-1)", {"a": s("abc")}, ("error", "#VALUE!"))
case("LEFT", "=LEFT([a],2)", {"a": n("12345")}, ("string", "12"))
case("RIGHT", "=RIGHT([a],2)", {"a": s("Nguyễn")}, ("string", "ễn"))
case("RIGHT", "=RIGHT([a],5)", {"a": s("abc")}, ("string", "abc"))
case("RIGHT", "=RIGHT([a],3)", {"a": s("00123")}, ("string", "123"))
case("MID", "=MID([a],2,3)", {"a": s("Nguyễn")}, ("string", "guy"))
case("MID", "=MID([a],5,2)", {"a": s("abc")}, ("blank", None), excel=("string", ""),
     note="Engine: chuỗi rỗng là ô trống.")
case("MID", "=MID([a],0,1)", {"a": s("abc")}, ("error", "#VALUE!"))
case("LEN", "=LEN([a])", {"a": s("Nguyễn")}, ("number", "6"))
case("LEN", "=LEN([a])", {"a": BLANK_S}, ("number", "0"))
case("LEN", "=LEN([a])", {"a": n("12.5")}, ("number", "4"))
case("TRIM", "=TRIM([a])", {"a": s("  Trần   Văn  Bình  ")}, ("string", "Trần Văn Bình"))
case("TRIM", "=TRIM([a])", {"a": s("abc")}, ("string", "abc"))
case("TRIM", '=TRIM([a])&"|"', {"a": BLANK_S}, ("string", "|"))
case("UPPER", "=UPPER([a])", {"a": s("nguyễn văn đức")}, ("string", "NGUYỄN VĂN ĐỨC"))
case("UPPER", "=UPPER([a])", {"a": n("1.5")}, ("string", "1.5"))
case("UPPER", '=UPPER([a])&"!"', {"a": BLANK_S}, ("string", "!"))
case("LOWER", "=LOWER([a])", {"a": s("ĐẠI HỌC")}, ("string", "đại học"))
case("LOWER", "=LOWER([a])", {"a": s("AbC")}, ("string", "abc"))
case("LOWER", "=LOWER([a])", {"a": b(True)}, ("string", "true"))
case("PROPER", "=PROPER([a])", {"a": s("nguyễn văn a")}, ("string", "Nguyễn Văn A"))
case("PROPER", "=PROPER([a])", {"a": s("o'neil-SMITH")}, ("string", "O'Neil-Smith"))
case("PROPER", "=PROPER([a])", {"a": s("abc123def")}, ("string", "Abc123Def"))
case("CONCAT", '=CONCAT("a",1,TRUE)', {}, ("string", "a1TRUE"))
case("CONCAT", '=CONCAT([a],"x")', {"a": BLANK_S}, ("string", "x"))
case("CONCAT", '=CONCAT([a]," ",[b])', {"a": s("Nguyễn"), "b": s("An")}, ("string", "Nguyễn An"))
case("TEXTJOIN", '=TEXTJOIN(", ",TRUE,"a","","b")', {}, ("string", "a, b"))
case("TEXTJOIN", '=TEXTJOIN(", ",FALSE,"a","","b")', {}, ("string", "a, , b"))
case("TEXTJOIN", '=TEXTJOIN("-",TRUE,[a],[b])', {"a": BLANK_S, "b": s("x")}, ("string", "x"))
case("SUBSTITUTE", '=SUBSTITUTE([a],"-","+")', {"a": s("a-b-c")}, ("string", "a+b+c"))
case("SUBSTITUTE", '=SUBSTITUTE([a],"-","+",2)', {"a": s("a-b-c")}, ("string", "a-b+c"))
case("SUBSTITUTE", '=SUBSTITUTE([a],"","x")', {"a": s("abc")}, ("string", "abc"))
case("SUBSTITUTE", '=SUBSTITUTE([a],"a","")', {"a": s("Aaa")}, ("string", "A"))
case("FIND", '=FIND("a","ABa")', {}, ("number", "3"))
case("FIND", '=FIND("z",[a])', {"a": s("abc")}, ("error", "#VALUE!"))
case("FIND", '=FIND("b",[a],3)', {"a": s("abcb")}, ("number", "4"))
case("FIND", '=FIND("ễ",[a])', {"a": s("Nguyễn")}, ("number", "5"))
case("SEARCH", '=SEARCH("b*","aBc")', {}, ("number", "2"))
case("SEARCH", '=SEARCH("?c",[a])', {"a": s("abc")}, ("number", "2"))
case("SEARCH", '=SEARCH("~*",[a])', {"a": s("a*b")}, ("number", "2"))
case("SEARCH", '=SEARCH("x",[a])', {"a": s("abc")}, ("error", "#VALUE!"))
case("SEARCH", '=SEARCH("VĂN",[a])', {"a": s("Nguyễn văn A")}, ("number", "8"))
case("VALUE", '=VALUE("1.5")', {}, ("number", "1.5"))
case("VALUE", "=VALUE([a])", {"a": s(" -12 ")}, ("number", "-12"))
case("VALUE", "=VALUE([a])", {"a": s("abc")}, ("error", "#VALUE!"))
case("VALUE", "=VALUE([a])", {"a": BLANK_S}, ("number", "0"))

# ---- Ngày ----------------------------------------------------------------------------------
case("DATE", "=DATE(2026,13,1)", {}, ("date", "2027-01-01"))
case("DATE", "=DATE(2024,2,29)", {}, ("date", "2024-02-29"))
case("DATE", "=DATE(2026,3,0)", {}, ("date", "2026-02-28"))
case("DATE", "=DATE(2026,1,-1)", {}, ("date", "2025-12-30"))
case("DATE", "=DATE(10000,1,1)", {}, ("error", "#NUM!"))
case("YEAR", "=YEAR([a])", {"a": d("2024-02-29")}, ("number", "2024"))
case("YEAR", "=YEAR([a])", {"a": d("2025-12-31")}, ("number", "2025"))
case("YEAR", "=YEAR([a])", {"a": BLANK_N}, ("number", "1900"), confirm="YEAR(ô trống)=1900")
case("YEAR", "=YEAR([a])", {"a": s("abc")}, ("error", "#VALUE!"))
case("MONTH", "=MONTH([a])", {"a": d("2024-02-29")}, ("number", "2"))
case("MONTH", "=MONTH([a])", {"a": d("2025-12-31")}, ("number", "12"))
case("MONTH", "=MONTH(-1)", {}, ("error", "#NUM!"))
case("DAY", "=DAY([a])", {"a": d("2024-02-29")}, ("number", "29"))
case("DAY", "=DAY([a])", {"a": d("2025-12-31")}, ("number", "31"))
case("DAY", "=DAY([a])", {"a": BLANK_N}, ("number", "0"), confirm="DAY(ô trống)=0")
case("EDATE", "=EDATE(DATE(2026,1,31),1)", {}, ("date", "2026-02-28"))
case("EDATE", "=EDATE([a],12)", {"a": d("2024-02-29")}, ("date", "2025-02-28"))
case("EDATE", "=EDATE([a],-1)", {"a": d("2026-03-31")}, ("date", "2026-02-28"))
case("EOMONTH", "=EOMONTH(DATE(2026,2,10),0)", {}, ("date", "2026-02-28"))
case("EOMONTH", "=EOMONTH([a],0)", {"a": d("2024-02-01")}, ("date", "2024-02-29"))
case("EOMONTH", "=EOMONTH([a],-2)", {"a": d("2026-01-15")}, ("date", "2025-11-30"))
case("DATEDIF", '=DATEDIF(DATE(2020,3,15),DATE(2026,3,14),"Y")', {}, ("number", "5"))
case("DATEDIF", '=DATEDIF([a],[b],"M")', {"a": d("2026-01-31"), "b": d("2026-02-28")}, ("number", "0"))
case("DATEDIF", '=DATEDIF([a],[b],"D")', {"a": d("2024-02-28"), "b": d("2024-03-01")}, ("number", "2"))
case("DATEDIF", '=DATEDIF([a],[b],"MD")', {"a": d("2026-01-31"), "b": d("2026-03-01")}, ("number", "-2"),
     confirm="DATEDIF MD khi ngày đầu lớn hơn ngày cuối tháng (lỗi đã biết của Excel)")
case("DATEDIF", '=DATEDIF([a],[b],"MD")', {"a": d("2026-01-15"), "b": d("2026-03-10")}, ("number", "23"),
     confirm="DATEDIF MD")
case("DATEDIF", '=DATEDIF([a],[b],"YM")', {"a": d("2020-11-15"), "b": d("2026-03-14")}, ("number", "3"))
case("DATEDIF", '=DATEDIF([a],[b],"YD")', {"a": d("2023-02-28"), "b": d("2024-03-01")}, ("number", "2"),
     confirm="DATEDIF YD qua năm nhuận")
case("DATEDIF", '=DATEDIF([a],[b],"Y")', {"a": d("2026-03-01"), "b": d("2026-01-01")}, ("error", "#NUM!"))
case("WEEKDAY", "=WEEKDAY(DATE(2026,10,30),2)", {}, ("number", "5"))
case("WEEKDAY", "=WEEKDAY([a])", {"a": d("2026-11-01")}, ("number", "1"))
case("WEEKDAY", "=WEEKDAY([a],3)", {"a": d("2024-02-29")}, ("number", "3"))
case("WEEKDAY", "=WEEKDAY([a],9)", {"a": d("2024-02-29")}, ("error", "#NUM!"))

# PARAM không phải hàm Excel: chỉ có trong unit test (transform-formula.test.ts).

for i, c in enumerate(CASES, start=1):
    c["id"] = f"c{i:03d}"

COLS = {"a": "D", "b": "E", "c": "F"}


def excel_formula(formula: str, row: int) -> str:
    out = re.sub(r"\[([abc])\]", lambda m: f"{COLS[m.group(1)]}{row}", formula)
    for fn in XLFN:
        out = re.sub(rf"(?<![A-Z_.]){fn}\(", f"_xlfn.{fn}(", out)
    return out


def cell_value(kind, v):
    if v is None:
        return None
    if kind == "n":
        return float(v) if ("." in v or "e" in v.lower()) else int(v)
    if kind == "d":
        return dt.date.fromisoformat(v)
    if kind == "b":
        return bool(v)
    return v


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "cases"
    header = ["id", "fn", "formula", "a", "b", "c", "excel", "error_type", "note"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row, c in enumerate(CASES, start=2):
        ws.cell(row=row, column=1, value=c["id"])
        ws.cell(row=row, column=2, value=c["fn"])
        # Công thức engine giữ dạng text (tiền tố ' để Excel không hiểu là công thức).
        ws.cell(row=row, column=3, value=c["formula"]).data_type = "s"
        ws.cell(row=row, column=3).quotePrefix = True
        for key, col in (("a", 4), ("b", 5), ("c", 6)):
            if key in c["inputs"]:
                kind, v = c["inputs"][key]
                cell = ws.cell(row=row, column=col, value=cell_value(kind, v))
                if kind == "d" and v is not None:
                    cell.number_format = "yyyy-mm-dd"
                if kind == "s" and v is not None:
                    cell.number_format = "@"
        ws.cell(row=row, column=7, value=excel_formula(c["formula"], row))
        ws.cell(row=row, column=8, value=f'=IF(ISERROR(G{row}),ERROR.TYPE(G{row}),"")')
        ws.cell(row=row, column=9, value=c.get("note") or c.get("confirm") or "")
    ws.column_dimensions["C"].width = 48
    ws.column_dimensions["G"].width = 22
    ws.column_dimensions["I"].width = 60
    wb.save(OUT / "golden.xlsx")

    out = []
    for c in CASES:
        item = {
            "id": c["id"],
            "fn": c["fn"],
            "formula": c["formula"],
            "inputs": {k: {"type": v[0], "value": v[1]} for k, v in c["inputs"].items()},
            "expect": {"type": c["expect"][0], "value": c["expect"][1]},
        }
        if c["expect"][0] == "date" and "excel" not in c:
            # Excel trả về ngày dưới dạng số serial (1900 date system).
            serial = (dt.date.fromisoformat(c["expect"][1]) - dt.date(1899, 12, 30)).days
            c["excel"] = ("number", str(serial))
            c["note"] = "Excel lưu ngày dạng số serial."
        if "excel" in c:
            item["excel"] = {"type": c["excel"][0], "value": c["excel"][1]}
            item["note"] = c.get("note", "")
        if "confirm" in c:
            item["confirm"] = c["confirm"]
        out.append(item)
    (OUT / "cases.json").write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"{len(CASES)} cases")


if __name__ == "__main__":
    main()
